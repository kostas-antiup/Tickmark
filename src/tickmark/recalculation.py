"""Recalculation engines that turn a workbook's formulas into values.

A workbook written by openpyxl stores formulas but no computed results, so value
checks need an engine that actually calculates:

- ``ExcelRecalculator`` drives Microsoft Excel through COM (Windows + pywin32);
- ``LibreOfficeRecalculator`` runs headless LibreOffice (any OS with ``soffice``);
- ``CachedValues`` only reads the values the file last saved and cannot apply
  input changes.
"""

from __future__ import annotations

import datetime
import importlib
import importlib.util
import os
import shutil
import subprocess
import sys
import tempfile
import warnings
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from concurrent.futures import ThreadPoolExecutor
from contextlib import ExitStack, suppress
from dataclasses import dataclass
from pathlib import Path
from types import TracebackType
from typing import Any, Protocol
from xml.etree import ElementTree

import openpyxl
from openpyxl.utils.cell import column_index_from_string, coordinate_from_string, get_column_letter

from .cell_refs import split_reference
from .xlsx_fast import XlsxTemplate, read_saved_values


class Recalculator(Protocol):
    """Evaluate cells of a workbook, optionally after overriding some input values."""

    name: str
    supports_changes: bool

    def evaluate(
        self,
        workbook: Path,
        cells: Iterable[str],
        changes: Mapping[str, Any] | None = None,
    ) -> dict[str, Any]:
        """Return ``{cell_key: value}`` for ``cells``; never modifies ``workbook``."""
        ...


class CachedValues:
    """Read the values the workbook last saved; no recalculation happens."""

    name = "cached"
    supports_changes = False

    def evaluate(
        self,
        workbook: Path,
        cells: Iterable[str],
        changes: Mapping[str, Any] | None = None,
    ) -> dict[str, Any]:
        if changes:
            raise ValueError("CachedValues cannot apply input changes; use a recalculating engine")
        return self._read(workbook, cells)

    def evaluate_many(
        self,
        workbook: Path,
        cells: Sequence[str],
        change_sets: Sequence[Mapping[str, Any]],
    ) -> list[dict[str, Any]]:
        if any(change_sets):
            raise ValueError("CachedValues cannot apply input changes; use a recalculating engine")
        values = self._read(workbook, cells)
        return [dict(values) for _ in change_sets]

    @staticmethod
    def _read(workbook: Path, cells: Iterable[str]) -> dict[str, Any]:
        book = openpyxl.load_workbook(workbook, data_only=True)
        values: dict[str, Any] = {}
        for key in cells:
            sheet, address = split_reference(key)
            values[key] = book[sheet][address].value if sheet in book.sheetnames else None
        return values


# Excel returns cell errors over COM as these integers (CVErr codes).
_EXCEL_ERRORS = {
    -2146826281: "#DIV/0!",
    -2146826246: "#N/A",
    -2146826259: "#NAME?",
    -2146826288: "#NULL!",
    -2146826252: "#NUM!",
    -2146826265: "#REF!",
    -2146826273: "#VALUE!",
}


class ExcelRecalculator:
    """Recalculate with a private, hidden Microsoft Excel instance via COM.

    Each call opens a temporary copy of the workbook, applies ``changes``, runs a
    full recalculation, reads the requested cells and closes without saving. Use
    as a context manager so the Excel process is shut down afterwards.
    """

    name = "excel"
    supports_changes = True

    def __init__(self) -> None:
        self._app: Any = None

    def __enter__(self) -> ExcelRecalculator:
        return self

    def __exit__(
        self,
        exc_type: type[BaseException] | None,
        exc: BaseException | None,
        traceback: TracebackType | None,
    ) -> None:
        self.close()

    def close(self) -> None:
        if self._app is not None:
            self._app.Quit()
            self._app = None

    def _excel(self) -> Any:
        if self._app is None:
            try:
                client = importlib.import_module("win32com.client")  # pywin32, Windows only
            except ImportError as error:
                raise RuntimeError(
                    "The excel engine needs Microsoft Excel and pywin32 on Windows "
                    "(`uv sync --extra audit`); use the libreoffice engine elsewhere."
                ) from error
            app = client.DispatchEx("Excel.Application")
            app.Visible = False
            app.DisplayAlerts = False
            app.AskToUpdateLinks = False
            app.EnableEvents = False
            self._app = app
        return self._app

    def evaluate(
        self,
        workbook: Path,
        cells: Iterable[str],
        changes: Mapping[str, Any] | None = None,
    ) -> dict[str, Any]:
        return self.evaluate_many(workbook, list(cells), [changes or {}])[0]

    def evaluate_many(
        self,
        workbook: Path,
        cells: Sequence[str],
        change_sets: Sequence[Mapping[str, Any]],
    ) -> list[dict[str, Any]]:
        """Open the workbook once; for each change set apply, recalculate, read, restore.

        Excel's COM server can die or drop the connection during a long run. On a COM
        error the instance is replaced by a fresh one and the call is retried once.
        """

        try:
            return self._evaluate_many(workbook, cells, change_sets)
        except _com_errors():
            self._restart()
            return self._evaluate_many(workbook, cells, change_sets)

    def _restart(self) -> None:
        with suppress(Exception):  # the old instance may already be gone
            self.close()
        self._app = None

    def _evaluate_many(
        self,
        workbook: Path,
        cells: Sequence[str],
        change_sets: Sequence[Mapping[str, Any]],
    ) -> list[dict[str, Any]]:
        app = self._excel()
        scratch = Path(tempfile.mkdtemp(prefix="fb_recalc_"))
        copy = scratch / Path(workbook).name
        shutil.copy2(workbook, copy)
        try:
            book = app.Workbooks.Open(str(copy), UpdateLinks=0, ReadOnly=True)
            try:
                sheets = {sheet.Name: sheet for sheet in book.Worksheets}
                blocks = _sheet_blocks(cells)
                results = []
                for changes in change_sets:
                    originals = {}
                    for key, value in changes.items():
                        sheet, address = split_reference(key)
                        if sheet not in sheets:  # a submission without the input sheet
                            continue
                        target = sheets[sheet].Range(address)
                        originals[key] = target.Formula
                        target.Value = value
                    app.CalculateFull()
                    values: dict[str, Any] = dict.fromkeys(cells)
                    for sheet, block in blocks.items():
                        if sheet in sheets:
                            values.update(_read_block(sheets[sheet], block))
                    results.append(values)
                    for key, formula in originals.items():
                        sheet, address = split_reference(key)
                        sheets[sheet].Range(address).Formula = formula
                return results
            finally:
                with suppress(Exception):  # closing fails too when Excel itself went away
                    book.Close(SaveChanges=False)
        finally:
            shutil.rmtree(scratch, ignore_errors=True)


def _com_errors() -> tuple[type[BaseException], ...]:
    """pywin32's COM error type, when pywin32 is installed (else nothing is caught)."""

    try:
        return (importlib.import_module("pywintypes").com_error,)
    except ImportError:
        return ()


# Read a sheet's cells as one rectangle while it stays this small; larger spreads cell by cell.
_MAX_BLOCK_CELLS = 250_000


@dataclass(frozen=True)
class _Block:
    """The requested cells of one sheet and the rectangle that covers them."""

    top: int
    left: int
    bottom: int
    right: int
    cells: tuple[tuple[str, int, int], ...]  # (key, row, column)

    @property
    def address(self) -> str:
        return (
            f"{get_column_letter(self.left)}{self.top}:{get_column_letter(self.right)}{self.bottom}"
        )


def _sheet_blocks(cells: Sequence[str]) -> dict[str, _Block]:
    by_sheet: dict[str, list[tuple[str, int, int]]] = {}
    for key in cells:
        sheet, address = split_reference(key)
        column, row = coordinate_from_string(address.replace("$", ""))
        by_sheet.setdefault(sheet, []).append((key, row, column_index_from_string(column)))
    return {
        sheet: _Block(
            top=min(row for _, row, _ in entries),
            left=min(column for _, _, column in entries),
            bottom=max(row for _, row, _ in entries),
            right=max(column for _, _, column in entries),
            cells=tuple(entries),
        )
        for sheet, entries in by_sheet.items()
    }


def _excel_value(value: Any) -> Any:
    if isinstance(value, int):
        return _EXCEL_ERRORS.get(value, value)
    if isinstance(value, datetime.datetime) and value.tzinfo is not None:
        # COM dates arrive time-zone aware; openpyxl's are naive. Keep them comparable.
        return datetime.datetime(*value.timetuple()[:6], value.microsecond)
    return value


def _read_block(sheet: Any, block: _Block) -> dict[str, Any]:
    """Read one sheet's requested cells; one COM call for the covering rectangle.

    Reading cell by cell costs a COM round trip per cell, which dominates when
    many scenarios are evaluated on a large model.
    """

    area = (block.bottom - block.top + 1) * (block.right - block.left + 1)
    if area > _MAX_BLOCK_CELLS:
        return {
            key: _excel_value(sheet.Cells(row, column).Value) for key, row, column in block.cells
        }
    raw = sheet.Range(block.address).Value
    grid = raw if isinstance(raw, tuple) else ((raw,),)
    return {
        key: _excel_value(grid[row - block.top][column - block.left])
        for key, row, column in block.cells
    }


def _write_variants(
    source: Path, targets: Sequence[Path], change_sets: Sequence[Mapping[str, Any]]
) -> None:
    """Save one copy of ``source`` per change set (formulas kept, values overridden).

    Input constants are patched straight into the sheet XML; a change set that needs
    more (a formula cell, a missing cell, text) is written with openpyxl instead.
    Changes on sheets the workbook lacks are skipped, so the copy simply ignores them.
    """

    try:
        template: XlsxTemplate | None = XlsxTemplate(source)
    except KeyError, zipfile.BadZipFile, ElementTree.ParseError:
        template = None
    book = None
    for target, changes in zip(targets, change_sets, strict=True):
        if not changes:
            shutil.copy2(source, target)
            continue
        if template is not None and template.write(target, changes):
            continue
        if book is None:
            book = openpyxl.load_workbook(source)
        originals = {}
        for key, value in changes.items():
            sheet, address = split_reference(key)
            if sheet not in book.sheetnames:
                continue
            originals[key] = book[sheet][address].value
            book[sheet][address] = value
        book.save(target)
        for key, value in originals.items():
            sheet, address = split_reference(key)
            book[sheet][address] = value


def _read_converted(workbook: Path, cells: Sequence[str]) -> dict[str, Any]:
    """Values LibreOffice saved; the XML reader first, openpyxl if the file surprises it."""

    try:
        return read_saved_values(workbook, cells)
    except Exception:
        return CachedValues._read(workbook, cells)


# Registry override for a private LibreOffice profile: always recalculate formulas when
# loading OOXML/ODF files instead of trusting (or missing) cached results.
_LIBREOFFICE_ALWAYS_RECALCULATE = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry"
  xmlns:xs="http://www.w3.org/2001/XMLSchema"
  xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load">
  <prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>
</item>
<item oor:path="/org.openoffice.Office.Calc/Formula/Load">
  <prop oor:name="ODFRecalcMode" oor:op="fuse"><value>0</value></prop>
</item>
</oor:items>
"""

_SOFFICE_CANDIDATES = (
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    "/Applications/LibreOffice.app/Contents/MacOS/soffice",
)


def find_soffice() -> str | None:
    """Locate the LibreOffice ``soffice`` executable, if installed."""

    for name in ("soffice", "libreoffice"):
        found = shutil.which(name)
        if found:
            return found
    return next((path for path in _SOFFICE_CANDIDATES if Path(path).exists()), None)


# One soffice process per this many files at least: starting LibreOffice costs about as
# much as converting a few small workbooks, so tiny batches stay in one process.
_MIN_FILES_PER_PROCESS = 3


def _default_workers() -> int:
    """Parallel soffice processes: TICKMARK_LIBREOFFICE_WORKERS, else up to 4 by CPU count."""

    configured = os.environ.get("TICKMARK_LIBREOFFICE_WORKERS", "")
    if configured.isdigit() and int(configured) > 0:
        return int(configured)
    return max(1, min(4, os.cpu_count() or 1))


class LibreOfficeRecalculator:
    """Recalculate with headless LibreOffice using private, throwaway profiles.

    Each call writes one temporary copy of the workbook per change set and converts
    them back to .xlsx; the profiles force a full recalculation on load, so the
    converted files hold fresh values, which are then read back. Larger batches are
    split over several soffice processes (each needs its own profile) running at once.
    """

    name = "libreoffice"
    supports_changes = True

    def __init__(
        self, soffice: str | None = None, timeout: float = 180.0, workers: int | None = None
    ) -> None:
        executable = soffice or find_soffice()
        if executable is None:
            raise RuntimeError(
                "LibreOffice (soffice) not found; install it or pass its path explicitly."
            )
        self._soffice = executable
        self._timeout = timeout
        self._workers = max(1, workers or _default_workers())
        self._profiles: dict[int, Path] = {}
        self._warm = False  # the first profile has been through LibreOffice's first start

    def __enter__(self) -> LibreOfficeRecalculator:
        return self

    def __exit__(
        self,
        exc_type: type[BaseException] | None,
        exc: BaseException | None,
        traceback: TracebackType | None,
    ) -> None:
        self.close()

    def close(self) -> None:
        for profile in self._profiles.values():
            shutil.rmtree(profile, ignore_errors=True)
        self._profiles.clear()
        self._warm = False

    def _profile_dir(self, index: int = 0) -> Path:
        """Profile ``index``; extra profiles are copies of the first once it is set up."""

        if index not in self._profiles:
            profile = Path(tempfile.mkdtemp(prefix="fb_lo_profile_"))
            if index and self._warm:
                shutil.copytree(
                    self._profiles[0],
                    profile,
                    dirs_exist_ok=True,
                    ignore=shutil.ignore_patterns(".lock"),
                )
            else:
                (profile / "user").mkdir()
                (profile / "user" / "registrymodifications.xcu").write_text(
                    _LIBREOFFICE_ALWAYS_RECALCULATE, encoding="utf-8"
                )
            self._profiles[index] = profile
        return self._profiles[index]

    def _warm_up(self, chunk: Sequence[Path], out_dir: Path) -> str:
        """Convert ``chunk`` alone on the first profile, then copy it for the other workers.

        LibreOffice's first start in a fresh profile does one-off setup; several such
        starts at once occasionally crash, so only one ever runs.
        """

        output = self._convert(0, chunk, out_dir)
        self._warm = True
        for index in range(1, self._workers):
            self._profile_dir(index)
        return output

    def _convert(self, index: int, sources: Sequence[Path], out_dir: Path) -> str:
        """Convert ``sources`` in one soffice process on profile ``index``; return its output."""

        command = [
            self._soffice,
            f"-env:UserInstallation={self._profile_dir(index).as_uri()}",
            "--headless",
            "--norestore",
            "--nologo",
            "--convert-to",
            "xlsx",
            "--outdir",
            str(out_dir),
            *map(str, sources),
        ]
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=self._timeout + 20 * len(sources),
            check=False,
        )
        return completed.stderr.strip() or completed.stdout.strip()

    def evaluate(
        self,
        workbook: Path,
        cells: Iterable[str],
        changes: Mapping[str, Any] | None = None,
    ) -> dict[str, Any]:
        return self.evaluate_many(workbook, list(cells), [changes or {}])[0]

    def evaluate_many(
        self,
        workbook: Path,
        cells: Sequence[str],
        change_sets: Sequence[Mapping[str, Any]],
    ) -> list[dict[str, Any]]:
        """Write one copy per change set and recalculate them all in as few soffice runs.

        Batches of at least ``_MIN_FILES_PER_PROCESS`` files per process are spread over
        up to ``workers`` processes that run at the same time.
        """

        if not change_sets:
            return []
        scratch = Path(tempfile.mkdtemp(prefix="fb_recalc_"))
        try:
            in_dir, out_dir = scratch / "in", scratch / "out"
            in_dir.mkdir()
            sources = [in_dir / f"scenario_{index:03d}.xlsx" for index in range(len(change_sets))]
            _write_variants(Path(workbook), sources, change_sets)
            count = min(self._workers, max(1, len(sources) // _MIN_FILES_PER_PROCESS))
            chunks = [sources[index::count] for index in range(count)]
            outputs = [] if self._warm else [self._warm_up(chunks[0], out_dir)]
            pending = list(range(len(outputs), count))
            if pending:
                with ThreadPoolExecutor(max_workers=len(pending)) as pool:
                    outputs += pool.map(
                        self._convert,
                        pending,
                        [chunks[index] for index in pending],
                        [out_dir] * len(pending),
                    )
            messages = {
                source: output for chunk, output in zip(chunks, outputs) for source in chunk
            }
            missing = [source for source in sources if not (out_dir / source.name).exists()]
            if missing:
                # A soffice process now and then dies mid-batch; its files get one more try.
                retry = self._convert(0, missing, out_dir)
                messages.update(dict.fromkeys(missing, retry))
            results = []
            for source in sources:
                converted = out_dir / source.name
                if not converted.exists():
                    raise RuntimeError(
                        f"LibreOffice could not recalculate {workbook}: {messages[source]}"
                    )
                results.append(_read_converted(converted, cells))
            return results
        finally:
            shutil.rmtree(scratch, ignore_errors=True)


def available_engine() -> str:
    """Best engine on this machine: Excel on Windows with pywin32, else LibreOffice, else cached."""

    if sys.platform == "win32" and importlib.util.find_spec("win32com") is not None:
        return "excel"
    if find_soffice() is not None:
        return "libreoffice"
    return "cached"


NO_ENGINE_WARNING = (
    "No Excel or LibreOffice found, so Tickmark reads the values saved in each file and skips "
    "the input-change and perturbation checks. Workbooks saved by openpyxl have no saved "
    "values and will fail the value check. Install LibreOffice for full grading."
)


def open_engine(name: str, stack: ExitStack) -> Recalculator:
    """Create the engine called ``name`` (``auto`` picks the best available) on ``stack``.

    ``auto`` warns when it has to fall back to cached values.
    """

    if name == "auto":
        name = available_engine()
        if name == "cached":
            warnings.warn(NO_ENGINE_WARNING, stacklevel=2)
    if name == "excel":
        return stack.enter_context(ExcelRecalculator())
    if name == "libreoffice":
        return stack.enter_context(LibreOfficeRecalculator())
    if name == "cached":
        return CachedValues()
    raise ValueError(f"unknown engine {name!r} (expected auto, excel, libreoffice or cached)")
