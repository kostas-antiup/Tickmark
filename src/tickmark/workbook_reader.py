"""Read an .xlsx file into sheet-qualified cell observations for grading."""

from __future__ import annotations

from collections.abc import Mapping
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl

from .cell_refs import DefinedNames, cell_key, expand_area, parse_formula
from .workbook_audit import CellObservation

UNRESOLVED_PREFIX = "?"


@dataclass(frozen=True)
class WorkbookSnapshot:
    """Formulas, cached values and references of every non-empty cell."""

    path: Path
    sheet_names: tuple[str, ...]
    observations: Mapping[str, CellObservation]
    literals: Mapping[str, tuple[float, ...]]
    issues: tuple[str, ...]


def formula_text(value: Any) -> str | None:
    """Return ``=...`` text for plain and array formulas, otherwise ``None``."""

    text = getattr(value, "text", value)  # openpyxl ArrayFormula keeps its formula in .text
    if isinstance(text, str) and text.startswith("="):
        return text
    return None


def read_workbook(path: Path | str) -> WorkbookSnapshot:
    """Load formulas and cached values from ``path``.

    Values are whatever the file last saved; a workbook written by openpyxl has
    no cached formula results until it is recalculated (see ``recalculation``).
    References that cannot be resolved inside the workbook (external links,
    unknown names) become precedents prefixed with ``?`` so tracing breaks there.
    """

    path = Path(path)
    formulas_book = openpyxl.load_workbook(path)
    values_book = openpyxl.load_workbook(path, data_only=True)
    bounds = {ws.title: (ws.max_row, ws.max_column) for ws in formulas_book.worksheets}
    names = _defined_names(formulas_book)

    observations: dict[str, CellObservation] = {}
    literals: dict[str, tuple[float, ...]] = {}
    issues: list[str] = []
    for sheet in formulas_book.worksheets:
        cached = values_book[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                raw = cell.value
                if raw is None or raw == "":
                    continue
                key = cell_key(sheet.title, cell.coordinate)
                formula = formula_text(raw)
                if formula is None:
                    observations.setdefault(key, CellObservation(key, raw))
                    continue
                refs = parse_formula(formula, sheet.title, sheet_bounds=bounds, defined_names=names)
                unresolved = {f"{UNRESOLVED_PREFIX}{token}" for token in refs.unresolved}
                if refs.unresolved:
                    issues.append(f"{key}: unresolved reference(s) {', '.join(refs.unresolved)}")
                for member in _array_members(sheet.title, raw, key):
                    member_address = member.split("!", 1)[1]
                    observations[member] = CellObservation(
                        member,
                        cached[member_address].value,
                        formula,
                        refs.direct | unresolved,
                        refs.ranged,
                    )
                    literals[member] = refs.literals
    return WorkbookSnapshot(
        path=path,
        sheet_names=tuple(formulas_book.sheetnames),
        observations=observations,
        literals=literals,
        issues=tuple(issues),
    )


def _array_members(sheet: str, raw: Any, key: str) -> list[str]:
    """Cells covered by a multi-cell array formula (openpyxl stores it on the first cell)."""

    ref = getattr(raw, "ref", None)
    if not isinstance(ref, str) or ":" not in ref:
        return [key]
    return expand_area(sheet, ref)


def _defined_names(book: Any) -> DefinedNames:
    names: dict[str, list[tuple[str, str]]] = {}
    scopes = [book.defined_names] + [
        sheet.defined_names for sheet in book.worksheets if hasattr(sheet, "defined_names")
    ]
    for scope in scopes:
        for name, definition in _iter_defined_names(scope):
            try:
                names[name.upper()] = list(definition.destinations)
            except AttributeError, ValueError:
                continue
    return names


def _iter_defined_names(scope: Any) -> list[tuple[str, Any]]:
    if hasattr(scope, "items"):
        return list(scope.items())
    return [(definition.name, definition) for definition in getattr(scope, "definedName", ())]
