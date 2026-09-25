"""Text view of a workbook for chat models, and applying the cell edits they return."""

from __future__ import annotations

import json
import re
from collections.abc import Mapping
from pathlib import Path
from typing import Any

import openpyxl

from .cell_refs import split_reference
from .workbook_reader import formula_text

MAX_EDITS = 5000
_ADDRESS = re.compile(r"^\$?[A-Z]{1,3}\$?\d+$")


def describe_workbook(path: Path) -> str:
    """One line per non-empty cell, sheet by sheet: ``C7: 58`` or ``G7: =SUM(C7:F7)  [= 320]``."""

    formulas = openpyxl.load_workbook(path)
    values = openpyxl.load_workbook(path, data_only=True)
    lines: list[str] = []
    for sheet in formulas.worksheets:
        lines.append(f'Sheet "{sheet.title}" (used range {sheet.dimensions}):')
        cached = values[sheet.title]
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None or cell.value == "":
                    continue
                formula = formula_text(cell.value)
                if formula is None:
                    lines.append(f"{cell.coordinate}: {cell.value!r}")
                    continue
                result = cached[cell.coordinate].value
                suffix = "" if result is None else f"  [= {result!r}]"
                lines.append(f"{cell.coordinate}: {formula}{suffix}")
        lines.append("")
    return "\n".join(lines)


def parse_cell_edits(reply: str, default_sheet: str) -> dict[str, Any]:
    """Parse ``{"cells": {"Sheet!A1": value, ...}}`` (or a flat mapping) from a model reply.

    Code fences and text around the JSON object are ignored. Addresses without a
    sheet name refer to ``default_sheet``. Raises ``ValueError`` with a reason the
    model can act on when the reply is unusable.
    """

    start, end = reply.find("{"), reply.rfind("}")
    if start == -1 or end <= start:
        raise ValueError("no JSON object found in the reply")
    try:
        data = json.loads(reply[start : end + 1])
    except json.JSONDecodeError as error:
        raise ValueError(f"invalid JSON ({error.msg} at position {error.pos})") from error
    cells = data.get("cells", data) if isinstance(data, dict) else None
    if not isinstance(cells, dict) or not cells:
        raise ValueError('expected a non-empty object {"cells": {"Sheet!A1": value}}')
    if len(cells) > MAX_EDITS:
        raise ValueError(f"too many cells ({len(cells)} > {MAX_EDITS})")

    edits: dict[str, Any] = {}
    for reference, value in cells.items():
        sheet, address = split_reference(str(reference), default_sheet)
        address = address.replace("$", "").upper()
        if not _ADDRESS.match(address):
            raise ValueError(f"{reference!r} is not a single-cell address like Sheet!C17")
        if not isinstance(value, (str, int, float)) or isinstance(value, bool):
            raise ValueError(f"value for {reference!r} must be a formula string or a number")
        edits[f"{sheet}!{address}"] = value
    return edits


def apply_cell_edits(source: Path, edits: Mapping[str, Any], target: Path) -> list[str]:
    """Write ``edits`` into a copy of ``source`` saved as ``target``; return skipped edits."""

    book = openpyxl.load_workbook(source)
    skipped: list[str] = []
    for key, value in edits.items():
        sheet, address = split_reference(key)
        if sheet not in book.sheetnames:
            skipped.append(f"{key}: sheet {sheet!r} does not exist")
            continue
        book[sheet][address] = value
    book.save(target)
    return skipped
