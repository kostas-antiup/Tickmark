"""Input scenarios for the behavioural checks.

Two formulas are treated as equivalent when they give the same values as the
reference model under many different inputs. This module builds those inputs:

- fixed scenarios: every input raised 5-25%, lowered 5-25%, mixed; all 0/1
  switches flipped; all year-like values shifted by one;
- branch-seeking scenarios: wider random inputs, kept only when they flip a
  decision (IF condition, MIN/MAX choice, IFERROR, ABS sign) in the reference
  model that the fixed scenarios never flipped; the share of decisions tested
  both ways is reported as branch coverage;
- one-at-a-time scenarios: each input changed alone, to name the inputs that
  make a submission diverge.

A ``ScenarioPlan`` (scenarios plus the reference model's results under each) depends
only on the case and the engine, so it is built once and shared by every
workbook graded for that case.
"""

from __future__ import annotations

import shutil
import tempfile
import zlib
from collections.abc import Iterable, Mapping, Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import TYPE_CHECKING, Any

import openpyxl
from openpyxl.formula.tokenizer import Tokenizer

from .cell_refs import split_reference

if TYPE_CHECKING:
    from .cases import CaseSpec
    from .recalculation import Recalculator
    from .workbook_audit import CellObservation

FACTOR_SCENARIOS = ("up", "down", "mixed")
DECISION_SHEET = "__fb_decisions"
SEARCH_CANDIDATES = 16
MAX_SEARCH_SCENARIOS = 10


@dataclass(frozen=True)
class Scenario:
    """Named set of input overrides (cell key -> new value)."""

    name: str
    changes: Mapping[str, Any]


@dataclass(frozen=True)
class Decision:
    """A branch point in a reference formula and a helper formula giving 1 or 0."""

    cell: str
    function: str
    helper: str

    @property
    def label(self) -> str:
        return f"{self.function} in {self.cell}"


def _numeric(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def is_switch(value: Any) -> bool:
    return _numeric(value) and value in (0, 1)


def is_year(value: Any) -> bool:
    return _numeric(value) and float(value).is_integer() and 1900 <= value <= 2100


def _scalable(value: Any) -> bool:
    return _numeric(value) and not is_switch(value) and not is_year(value)


def factor_scenario(inputs: Mapping[str, Any], name: str) -> Scenario:
    """``up``/``down``/``mixed``: each scalable input changed by its own 5-25%."""

    if name not in FACTOR_SCENARIOS:
        raise ValueError(f"unknown scenario {name!r}; expected one of {FACTOR_SCENARIOS}")
    changes: dict[str, float] = {}
    for key, value in sorted(inputs.items()):
        if not _scalable(value):
            continue
        digest = zlib.crc32(f"{name}:{key}".encode())
        step = 0.05 + (digest % 21) / 100
        rises = name == "up" or (name == "mixed" and digest % 2 == 0)
        changes[key] = value * (1 + step if rises else 1 - step)
    return Scenario(name, changes)


def fixed_scenarios(inputs: Mapping[str, Any]) -> list[Scenario]:
    """Up, down, mixed, plus switch flips and year shifts when the model has them."""

    scenarios = [factor_scenario(inputs, name) for name in FACTOR_SCENARIOS]
    switches = {key: 1 - value for key, value in sorted(inputs.items()) if is_switch(value)}
    years = {key: value + 1 for key, value in sorted(inputs.items()) if is_year(value)}
    if switches:
        scenarios.append(Scenario("switches", switches))
    if years:
        scenarios.append(Scenario("years", years))
    return [scenario for scenario in scenarios if scenario.changes]


def search_candidates(inputs: Mapping[str, Any], count: int = SEARCH_CANDIDATES) -> list[Scenario]:
    """Wider deterministic scenarios: inputs x0.4-1.9, switches flipped at random."""

    candidates = []
    for index in range(1, count + 1):
        changes: dict[str, Any] = {}
        for key, value in sorted(inputs.items()):
            digest = zlib.crc32(f"search{index}:{key}".encode())
            if _scalable(value):
                changes[key] = value * (0.4 + (digest % 151) / 100)
            elif is_switch(value) and digest % 2:
                changes[key] = 1 - value
        if changes:
            candidates.append(Scenario(f"branch-search-{index}", changes))
    return candidates


def extreme_candidates(inputs: Mapping[str, Any], relevant: Iterable[str]) -> list[Scenario]:
    """One input at a time pushed far: x4 and x0.25 (switches flipped alone)."""

    candidates = []
    for key in sorted(set(relevant) & set(inputs)):
        value = inputs[key]
        if _scalable(value):
            candidates.append(Scenario(f"branch-extreme-{key}-x4", {key: value * 4}))
            candidates.append(Scenario(f"branch-extreme-{key}-x0.25", {key: value * 0.25}))
        elif is_switch(value):
            candidates.append(Scenario(f"branch-extreme-{key}-flip", {key: 1 - value}))
    return candidates


def one_at_a_time(inputs: Mapping[str, Any]) -> list[Scenario]:
    """Each input changed alone: switches flipped, years +1, other values +15%."""

    scenarios = []
    for key, value in sorted(inputs.items()):
        if is_switch(value):
            new: Any = 1 - value
        elif is_year(value):
            new = value + 1
        elif _numeric(value):
            new = value * 1.15
        else:
            continue
        scenarios.append(Scenario(f"only:{key}", {key: new}))
    return scenarios


# --- decision points in reference formulas ------------------------------------------------


def _qualified(tokens: Sequence[Any], sheet: str) -> str:
    """Rebuild tokens as text, adding ``'sheet'!`` to unqualified cell references."""

    quoted = "'" + sheet.replace("'", "''") + "'!"
    parts = []
    for token in tokens:
        value = token.value
        if token.type == "OPERAND" and token.subtype == "RANGE" and "!" not in value:
            if any(ch.isdigit() for ch in value) or ":" in value:
                value = quoted + value
        parts.append(value)
    return "".join(parts)


def _arguments(tokens: Sequence[Any], open_index: int) -> list[list[Any]]:
    """Top-level argument token lists of the function opened at ``open_index``."""

    args: list[list[Any]] = [[]]
    depth = 0
    for token in tokens[open_index + 1 :]:
        opening = token.subtype == "OPEN" and token.type in ("FUNC", "PAREN")
        closing = token.subtype == "CLOSE" and token.type in ("FUNC", "PAREN")
        if closing and depth == 0:
            return args
        if opening:
            depth += 1
        elif closing:
            depth -= 1
        if depth == 0 and token.type == "SEP" and token.subtype == "ARG":
            args.append([])
            continue
        args[-1].append(token)
    return args


def _is_range(arg: Sequence[Any]) -> bool:
    meaningful = [t for t in arg if t.type != "WHITE-SPACE"]
    return len(meaningful) == 1 and meaningful[0].subtype == "RANGE" and ":" in meaningful[0].value


def extract_decisions(
    observations: Mapping[str, CellObservation], targets: Iterable[str]
) -> list[Decision]:
    """Branch points (IF, two-argument MIN/MAX, IFERROR, ABS) in the targets' formulas."""

    decisions = []
    for key in targets:
        observation = observations.get(key)
        if observation is None or not observation.has_formula or observation.formula is None:
            continue
        sheet = split_reference(key)[0]
        try:
            tokens = Tokenizer(observation.formula).items
        except Exception:
            continue
        for index, token in enumerate(tokens):
            if token.type != "FUNC" or token.subtype != "OPEN":
                continue
            name = token.value.rstrip("(").upper().replace("_XLFN.", "")
            args = _arguments(tokens, index)
            texts = [_qualified(arg, sheet) for arg in args]
            if name == "IF" and texts and texts[0].strip():
                helper = f"=IF(({texts[0]}),1,0)"
            elif name in ("MIN", "MAX") and len(args) == 2 and not any(map(_is_range, args)):
                helper = f"=IF(({texts[0]})<=({texts[1]}),1,0)"
            elif name == "IFERROR" and texts:
                helper = f"=IF(ISERROR({texts[0]}),1,0)"
            elif name == "ABS" and len(texts) == 1:
                helper = f"=IF(({texts[0]})>=0,1,0)"
            else:
                continue
            decisions.append(Decision(key, name, helper))
    return decisions


# --- test plan ------------------------------------------------------------------------------


@dataclass
class ScenarioPlan:
    """Scenarios to compare and the reference model's target values under each."""

    scenarios: list[Scenario]
    expected: dict[str, dict[str, Any]]
    decisions: list[Decision] = field(default_factory=list)
    covered: list[bool] = field(default_factory=list)
    one_at_a_time: list[Scenario] | None = None
    one_at_a_time_expected: dict[str, dict[str, Any]] | None = None

    @property
    def untested(self) -> list[Decision]:
        return [d for d, ok in zip(self.decisions, self.covered, strict=True) if not ok]


_PLANS: dict[tuple[str, int, str, str], ScenarioPlan] = {}


def evaluate_many(
    engine: Recalculator,
    workbook: Path,
    cells: Sequence[str],
    change_sets: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    """Use the engine's batch method when it has one, else one call per change set."""

    batch = getattr(engine, "evaluate_many", None)
    if batch is not None:
        return batch(workbook, cells, change_sets)
    return [engine.evaluate(workbook, cells, changes) for changes in change_sets]


def _outcome(value: Any) -> int | None:
    if _numeric(value) and value in (0, 1):
        return int(value)
    return None


def scenario_plan(case: CaseSpec, engine: Recalculator) -> ScenarioPlan:
    """Build (once per case and engine) the scenarios and reference results."""

    key = (engine.name, id(engine), case.case_id, str(case.golden_workbook))
    if key not in _PLANS:
        _PLANS[key] = _build_plan(case, engine)
    return _PLANS[key]


def clear_plans() -> None:
    _PLANS.clear()


def _build_plan(case: CaseSpec, engine: Recalculator) -> ScenarioPlan:
    inputs = case.golden_input_values
    targets = list(case.target_cells)
    fixed = fixed_scenarios(inputs)
    decisions = list(case.reference_decisions)
    if not decisions or not case.golden_workbook.exists():
        results = evaluate_many(engine, case.golden_workbook, targets, [s.changes for s in fixed])
        return ScenarioPlan(fixed, {s.name: r for s, r in zip(fixed, results, strict=True)})

    scratch = Path(tempfile.mkdtemp(prefix="fb_plan_"))
    try:
        helper_book = scratch / case.golden_workbook.name
        book = openpyxl.load_workbook(case.golden_workbook)
        sheet = book.create_sheet(DECISION_SHEET)
        helper_cells = []
        for row, decision in enumerate(decisions, start=1):
            sheet.cell(row=row, column=1, value=decision.helper)
            helper_cells.append(f"{DECISION_SHEET}!A{row}")
        book.save(helper_book)

        # Helper outcomes and target values come from the same recalculation, so each
        # phase below is a single engine run (one LibreOffice start).
        cells = [*helper_cells, *targets]
        seen: list[set[int]] = [set() for _ in decisions]
        expected: dict[str, dict[str, Any]] = {}
        kept: list[Scenario] = []

        def run(always: Sequence[Scenario], candidates: Sequence[Scenario]) -> None:
            """Record ``always``; keep a candidate only if it shows a new decision outcome."""

            scenarios = [*always, *candidates]
            results = evaluate_many(engine, helper_book, cells, [s.changes for s in scenarios])
            for position, (scenario, result) in enumerate(zip(scenarios, results, strict=True)):
                search = position >= len(always)
                if search and len(kept) >= MAX_SEARCH_SCENARIOS:
                    return
                new = [
                    (index, value)
                    for index, value in enumerate(_outcome(result[c]) for c in helper_cells)
                    if value is not None and value not in seen[index]
                ]
                for index, value in new:
                    seen[index].add(value)
                if search and not new:
                    continue
                if search:
                    kept.append(scenario)
                expected[scenario.name] = {t: result[t] for t in targets}

        # Phase 1: the fixed scenarios plus wide random ones. Phase 2, for decisions still
        # tested one way only: push each input they depend on alone to x4 and x0.25.
        run([Scenario("base", {}), *fixed], search_candidates(inputs))
        untested = [d for d, values in zip(decisions, seen, strict=True) if len(values) < 2]
        if untested and len(kept) < MAX_SEARCH_SCENARIOS:
            relevant = set().union(
                *(case.reference_footprints.get(d.cell, frozenset(inputs)) for d in untested)
            )
            candidates = extreme_candidates(inputs, relevant)
            if candidates:
                run([], candidates)
    finally:
        shutil.rmtree(scratch, ignore_errors=True)

    expected.pop("base", None)
    return ScenarioPlan([*fixed, *kept], expected, decisions, [len(s) == 2 for s in seen])


def one_at_a_time_expected(
    case: CaseSpec, engine: Recalculator, plan: ScenarioPlan
) -> tuple[list[Scenario], dict[str, dict[str, Any]]]:
    """Reference results with each input changed alone (built lazily, then cached)."""

    if plan.one_at_a_time is None or plan.one_at_a_time_expected is None:
        scenarios = one_at_a_time(case.golden_input_values)
        results = evaluate_many(
            engine, case.golden_workbook, list(case.target_cells), [s.changes for s in scenarios]
        )
        plan.one_at_a_time = scenarios
        plan.one_at_a_time_expected = {s.name: r for s, r in zip(scenarios, results, strict=True)}
    return plan.one_at_a_time, plan.one_at_a_time_expected
