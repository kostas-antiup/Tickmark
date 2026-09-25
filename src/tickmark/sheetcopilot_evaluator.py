"""Evaluate a workbook against SheetCopilot reference checklists."""

from __future__ import annotations

from collections.abc import Callable
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any

import openpyxl

from .sheetcopilot_cases import SheetCopilotCase, SheetCopilotReference
from .workbook_audit import values_match

DEFAULT_ABS_TOLERANCE = 1e-8
DEFAULT_REL_TOLERANCE = 1e-8

_SUPPORTED_CHECKS: dict[tuple[str, str], str] = {
    ("cells", "values"): "cell values",
    ("cells", "formatting"): "cell formatting",
    ("cells", "hyperlink"): "cell hyperlinks",
    ("filters", "filters"): "filter settings",
    ("filters", "source"): "filter source range",
    ("format_conditions", "bold"): "conditional formatting",
    ("format_conditions", "color"): "conditional formatting",
    ("format_conditions", "fill_color"): "conditional formatting",
    ("format_conditions", "font"): "conditional formatting",
    ("format_conditions", "formula1"): "conditional formatting",
    ("format_conditions", "italic"): "conditional formatting",
    ("format_conditions", "operator"): "conditional formatting",
    ("format_conditions", "type"): "conditional formatting",
    ("format_conditions", "underline"): "conditional formatting",
    ("view", "freeze_pane"): "frozen panes",
}


@dataclass(frozen=True)
class SheetCopilotCheckResult:
    """Result for one enabled SheetCopilot checklist field."""

    reference_id: str
    sheet_index: str
    aspect: str
    field: str
    passed: bool
    detail: str
    mismatches: tuple[str, ...] = ()


@dataclass(frozen=True)
class SheetCopilotReferenceResult:
    """All checks for one acceptable SheetCopilot reference workbook."""

    reference_id: str
    passed: bool
    checks: tuple[SheetCopilotCheckResult, ...]


@dataclass(frozen=True)
class SheetCopilotEvaluationReport:
    """JSON-ready evaluation report for one SheetCopilot case."""

    case_id: str
    workbook: str
    passed: bool
    matched_reference: str | None
    references: tuple[SheetCopilotReferenceResult, ...]
    required_apis: tuple[str, ...]
    remarks: tuple[str, ...]

    def to_dict(self) -> dict[str, Any]:
        return asdict(self)


def evaluate_sheetcopilot_workbook(
    case: SheetCopilotCase,
    workbook: Path | str,
    *,
    abs_tolerance: float = DEFAULT_ABS_TOLERANCE,
    rel_tolerance: float = DEFAULT_REL_TOLERANCE,
) -> SheetCopilotEvaluationReport:
    """Compare ``workbook`` with every reference accepted for ``case``."""

    workbook = Path(workbook)
    if not case.references:
        return SheetCopilotEvaluationReport(
            case.case_id,
            str(workbook),
            False,
            None,
            (),
            (),
            (f"No references found in {case.reference_dir}.",),
        )

    actual_values = openpyxl.load_workbook(workbook, data_only=True)
    actual_shapes = openpyxl.load_workbook(workbook)
    results = tuple(
        _evaluate_reference(
            reference,
            actual_values,
            actual_shapes,
            abs_tolerance=abs_tolerance,
            rel_tolerance=rel_tolerance,
        )
        for reference in case.references
    )
    matched = next((result.reference_id for result in results if result.passed), None)
    required_apis = tuple(
        dict.fromkeys(api for ref in case.references for api in ref.required_apis)
    )
    remarks: list[str] = []
    if matched is None:
        remarks.append(
            "Workbook did not match any SheetCopilot reference checklist. "
            f"Best reference: {_best_reference(results)}."
        )
    return SheetCopilotEvaluationReport(
        case.case_id,
        str(workbook),
        matched is not None,
        matched,
        results,
        required_apis,
        tuple(remarks),
    )


def _evaluate_reference(
    reference: SheetCopilotReference,
    actual_values: Any,
    actual_shapes: Any,
    *,
    abs_tolerance: float,
    rel_tolerance: float,
) -> SheetCopilotReferenceResult:
    reference_values = openpyxl.load_workbook(reference.workbook, data_only=True)
    reference_shapes = openpyxl.load_workbook(reference.workbook)
    checks: list[SheetCopilotCheckResult] = []
    for sheet_index, aspects in reference.check_board.items():
        for aspect, fields in aspects.items():
            for field, enabled in fields.items():
                if not enabled:
                    continue
                checks.append(
                    _run_check(
                        reference.reference_id,
                        sheet_index,
                        aspect,
                        field,
                        reference_values,
                        reference_shapes,
                        actual_values,
                        actual_shapes,
                        abs_tolerance=abs_tolerance,
                        rel_tolerance=rel_tolerance,
                    )
                )
    return SheetCopilotReferenceResult(
        reference.reference_id,
        bool(checks) and all(check.passed for check in checks),
        tuple(checks),
    )


def _run_check(
    reference_id: str,
    sheet_index: str,
    aspect: str,
    field: str,
    reference_values: Any,
    reference_shapes: Any,
    actual_values: Any,
    actual_shapes: Any,
    *,
    abs_tolerance: float,
    rel_tolerance: float,
) -> SheetCopilotCheckResult:
    label = _SUPPORTED_CHECKS.get((aspect, field))
    if label is None:
        return SheetCopilotCheckResult(
            reference_id,
            sheet_index,
            aspect,
            field,
            False,
            f"Unsupported SheetCopilot check: {aspect}.{field}.",
        )

    ref_value_sheet = _sheet_by_index(reference_values, sheet_index)
    ref_shape_sheet = _sheet_by_index(reference_shapes, sheet_index)
    actual_value_sheet = _sheet_by_index(actual_values, sheet_index)
    actual_shape_sheet = _sheet_by_index(actual_shapes, sheet_index)
    if (
        ref_value_sheet is None
        or ref_shape_sheet is None
        or actual_value_sheet is None
        or actual_shape_sheet is None
    ):
        return SheetCopilotCheckResult(
            reference_id,
            sheet_index,
            aspect,
            field,
            False,
            f"Missing sheet index {sheet_index}.",
        )

    handlers: dict[tuple[str, str], Callable[[], tuple[bool, str, tuple[str, ...]]]] = {
        ("cells", "values"): lambda: _compare_cells(
            ref_value_sheet,
            actual_value_sheet,
            lambda ref, actual: cell_values_match(
                ref, actual, abs_tolerance=abs_tolerance, rel_tolerance=rel_tolerance
            ),
            "value",
        ),
        ("cells", "formatting"): lambda: _compare_cells(
            ref_shape_sheet,
            actual_shape_sheet,
            lambda ref, actual: _cell_format(ref) == _cell_format(actual),
            "formatting",
        ),
        ("cells", "hyperlink"): lambda: _compare_cells(
            ref_shape_sheet,
            actual_shape_sheet,
            lambda ref, actual: _hyperlink(ref) == _hyperlink(actual),
            "hyperlink",
        ),
        ("view", "freeze_pane"): lambda: _compare_scalar(
            ref_shape_sheet.freeze_panes,
            actual_shape_sheet.freeze_panes,
            "freeze panes",
        ),
        ("filters", "source"): lambda: _compare_scalar(
            ref_shape_sheet.auto_filter.ref,
            actual_shape_sheet.auto_filter.ref,
            "auto-filter range",
        ),
        ("filters", "filters"): lambda: _compare_scalar(
            _filter_signature(ref_shape_sheet),
            _filter_signature(actual_shape_sheet),
            "auto-filter settings",
        ),
    }
    if aspect == "format_conditions":
        outcome = _compare_scalar(
            _conditional_formatting_signature(ref_shape_sheet),
            _conditional_formatting_signature(actual_shape_sheet),
            "conditional formatting",
        )
    else:
        outcome = handlers[(aspect, field)]()
    passed, detail, mismatches = outcome
    return SheetCopilotCheckResult(
        reference_id, sheet_index, aspect, field, passed, detail, mismatches
    )


def _sheet_by_index(book: Any, sheet_index: str) -> Any | None:
    try:
        index = int(sheet_index) - 1
    except ValueError:
        return None
    return book.worksheets[index] if 0 <= index < len(book.worksheets) else None


def _compare_cells(
    reference_sheet: Any,
    actual_sheet: Any,
    same: Callable[[Any, Any], bool],
    label: str,
) -> tuple[bool, str, tuple[str, ...]]:
    mismatches: list[str] = []
    checked = 0
    for row in range(1, reference_sheet.max_row + 1):
        for column in range(1, reference_sheet.max_column + 1):
            ref_cell = reference_sheet.cell(row, column)
            actual_cell = actual_sheet.cell(row, column)
            checked += 1
            if not same(ref_cell, actual_cell):
                mismatches.append(ref_cell.coordinate)
    detail = f"{checked - len(mismatches)}/{checked} checked cell {label}(s) match."
    return not mismatches, detail, tuple(mismatches)


def _compare_scalar(expected: Any, actual: Any, label: str) -> tuple[bool, str, tuple[str, ...]]:
    passed = expected == actual
    detail = f"{label}: expected {expected!r}, actual {actual!r}."
    return passed, detail, () if passed else (label,)


def cell_values_match(
    reference: Any,
    actual: Any,
    *,
    abs_tolerance: float = DEFAULT_ABS_TOLERANCE,
    rel_tolerance: float = DEFAULT_REL_TOLERANCE,
) -> bool:
    """SheetCopilot's ``cells.values`` rule for one cell (cells or plain values accepted)."""

    return values_match(
        _blank(reference), _blank(actual), abs_tolerance=abs_tolerance, rel_tolerance=rel_tolerance
    )


def _blank(value: Any) -> Any:
    value = getattr(value, "value", value)
    return None if value == "" else value


def _hyperlink(cell: Any) -> tuple[str | None, str | None]:
    link = cell.hyperlink
    if link is None:
        return (None, None)
    return (link.target, link.location)


def _cell_format(cell: Any) -> tuple[Any, ...]:
    return (
        cell.number_format,
        cell.font.name,
        cell.font.sz,
        cell.font.bold,
        cell.font.italic,
        cell.font.underline,
        cell.font.color.type if cell.font.color else None,
        cell.font.color.rgb if cell.font.color and cell.font.color.type == "rgb" else None,
        cell.fill.fill_type,
        cell.fill.fgColor.type,
        cell.fill.fgColor.rgb,
        cell.alignment.horizontal,
        cell.alignment.vertical,
        cell.protection.locked,
        cell.protection.hidden,
    )


def _filter_signature(sheet: Any) -> tuple[Any, ...]:
    auto_filter = sheet.auto_filter
    sort_state = auto_filter.sortState or ()
    return (
        auto_filter.ref,
        tuple(repr(column) for column in auto_filter.filterColumn),
        tuple(repr(item) for item in sort_state),
    )


def _conditional_formatting_signature(sheet: Any) -> tuple[Any, ...]:
    signature: list[Any] = []
    for conditional_formatting in sheet.conditional_formatting:
        rules = sheet.conditional_formatting[conditional_formatting]
        signature.append(
            (
                str(conditional_formatting),
                tuple(
                    (
                        rule.type,
                        rule.operator,
                        tuple(rule.formula or ()),
                        _style_signature(rule.dxf),
                    )
                    for rule in rules
                ),
            )
        )
    return tuple(signature)


def _style_signature(style: Any) -> tuple[Any, ...]:
    if style is None:
        return ()
    font = style.font
    fill = style.fill
    return (
        getattr(font, "b", None),
        getattr(font, "i", None),
        getattr(font, "u", None),
        getattr(getattr(font, "color", None), "rgb", None),
        getattr(fill, "fill_type", None),
        getattr(getattr(fill, "fgColor", None), "rgb", None),
    )


def _best_reference(results: tuple[SheetCopilotReferenceResult, ...]) -> str:
    if not results:
        return "none"
    best = max(results, key=lambda result: sum(check.passed for check in result.checks))
    passed = sum(check.passed for check in best.checks)
    return f"{best.reference_id} ({passed}/{len(best.checks)} checks passed)"
