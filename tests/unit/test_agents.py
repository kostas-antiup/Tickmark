import json
import sys
import threading
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import openpyxl
import pytest
from hamcrest import assert_that, contains_string, equal_to, has_length, none
from tests.shared.givenpy import given, then, when

from tickmark.agents import (
    OUTPUT_NAME,
    SHORT_PROMPT,
    TASK_NAME,
    AgentTask,
    ChatAgent,
    CliAgent,
    ManualAgent,
    check_agent,
    load_agent_table,
    load_agents,
    task_text,
)

CONFIG = Path(__file__).resolve().parents[2] / "configs" / "agents.toml"


def _task(tmp_path) -> AgentTask:
    book = openpyxl.Workbook()
    book.worksheets[0].title = "Model"
    book.worksheets[0]["B2"] = 100
    source = tmp_path / "source.xlsx"
    book.save(source)
    return AgentTask("toy", "Fill Model!B10 with =B2*2.", source, tmp_path / "workspace")


# A stand-in for a coding agent: checks TASK.md exists, echoes the prompt and writes
# output.xlsx (or edits input.xlsx in place) with openpyxl.
FAKE_AGENT = """
import json, sys, pathlib, openpyxl
mode, prompt = sys.argv[1], sys.argv[2]
assert pathlib.Path("TASK.md").exists()
if mode == "crash":
    print(json.dumps({"is_error": True, "result": "Not logged in", "echo": prompt}))
    sys.exit(1)
print(prompt)
if mode != "nothing":
    book = openpyxl.load_workbook("input.xlsx")
    book["Model"]["B10"] = "=B2*2"
    book.save("output.xlsx" if mode == "output" else "input.xlsx")
"""


@pytest.mark.parametrize(("mode", "status"), [("output", "completed"), ("nothing", "no_output")])
def test_when_cli_agent_runs_then_its_workbook_is_collected(tmp_path, mode, status) -> None:
    with given() as context:
        context.agent = CliAgent("fake", [sys.executable, "-c", FAKE_AGENT, mode, "{prompt}"])

    with when():
        context.result = context.agent.solve(_task(tmp_path))

    with then():
        assert_that(context.result.status, equal_to(status))
        assert_that(context.result.log["stdout_tail"], contains_string(SHORT_PROMPT))
        if status == "completed":
            output = openpyxl.load_workbook(context.result.output_workbook)["Model"]
            assert_that(output["B10"].value, equal_to("=B2*2"))


def test_when_cli_agent_edits_input_in_place_then_that_workbook_is_graded(tmp_path) -> None:
    with given() as context:
        context.agent = CliAgent("fake", [sys.executable, "-c", FAKE_AGENT, "inplace", "{prompt}"])

    with when():
        context.result = context.agent.solve(_task(tmp_path))

    with then():
        assert_that(context.result.status, equal_to("completed"))
        output = context.result.output_workbook
        assert output is not None
        assert_that(output.name, equal_to("input.xlsx"))


def test_when_cli_agent_exits_with_an_error_then_its_message_is_reported(tmp_path) -> None:
    with given() as context:
        context.agent = CliAgent("fake", [sys.executable, "-c", FAKE_AGENT, "crash", "{prompt}"])

    with when():
        context.result = context.agent.solve(_task(tmp_path))

    with then():
        assert_that(context.result.status, equal_to("failed"))
        assert_that(context.result.error, equal_to("agent failed: Not logged in"))
        assert_that(context.result.log["stdout_json"]["echo"], equal_to(SHORT_PROMPT))


def test_when_cli_executable_is_missing_then_the_run_fails_with_a_reason(tmp_path) -> None:
    with given() as context:
        context.agent = CliAgent("ghost", ["no-such-agent-binary", "{prompt}"])

    with when():
        context.result = context.agent.solve(_task(tmp_path))

    with then():
        assert_that(context.result.status, equal_to("failed"))
        assert_that(context.result.error or "", contains_string("not found"))


def test_when_cli_task_is_prepared_then_it_requires_formula_first_workbook_editing(
    tmp_path,
) -> None:
    with given() as context:
        context.task = _task(tmp_path)

    with when():
        context.text = task_text(context.task)

    with then():
        assert_that(context.text, contains_string("Use Python with `openpyxl`"))
        assert_that(context.text, contains_string("Do not unzip the\nworkbook"))
        assert_that(
            context.text,
            contains_string("Save `output.xlsx` before any\noptional checking"),
        )
        assert_that(context.text, contains_string("each requested target cell contains a formula"))


def _reply(content):
    return {"choices": [{"message": {"content": content}}], "usage": {"total_tokens": 42}}


def test_when_chat_reply_is_unusable_then_one_repair_turn_fixes_it(tmp_path, monkeypatch) -> None:
    with given() as context:
        monkeypatch.setenv("FAKE_KEY", "secret")
        context.calls = []

        def transport(url, payload, headers, timeout):
            context.calls.append((url, payload, headers))
            if len(context.calls) == 1:
                return _reply("Sure! B10 should double B2.")
            return _reply('{"cells": {"Model!B10": "=B2*2"}}')

        context.agent = ChatAgent(
            "fake-chat",
            base_url="https://example.test/v1/",
            model="m",
            api_key_env="FAKE_KEY",
            transport=transport,
        )

    with when():
        context.result = context.agent.solve(_task(tmp_path))

    with then():
        url, payload, headers = context.calls[-1]
        assert_that(url, equal_to("https://example.test/v1/chat/completions"))
        assert_that(headers["Authorization"], equal_to("Bearer secret"))
        assert_that(payload["messages"], has_length(4))
        assert_that(payload["messages"][1]["content"], contains_string("B2: 100"))
        assert_that(context.result.status, equal_to("completed"))
        assert_that(context.result.log["exchanges"], has_length(2))
        output = openpyxl.load_workbook(context.result.output_workbook)["Model"]
        assert_that(output["B10"].value, equal_to("=B2*2"))


def test_when_chat_api_key_is_missing_then_the_run_fails_with_a_reason(
    tmp_path, monkeypatch
) -> None:
    with given() as context:
        monkeypatch.delenv("MISSING_KEY", raising=False)
        context.agent = ChatAgent(
            "fake-chat", base_url="https://example.test/v1", model="m", api_key_env="MISSING_KEY"
        )

    with when():
        context.result = context.agent.solve(_task(tmp_path))

    with then():
        assert_that(context.result.status, equal_to("failed"))
        assert_that(context.result.error or "", contains_string("MISSING_KEY"))


class _CompletionsHandler(BaseHTTPRequestHandler):
    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        assert body["model"] == "local-model"
        reply = json.dumps(_reply('{"cells": {"Model!B10": "=B2*2"}}')).encode()
        self.send_response(200)
        self.send_header("Content-Type", "application/json")
        self.send_header("Content-Length", str(len(reply)))
        self.end_headers()
        self.wfile.write(reply)

    def log_message(self, format, *args):
        pass


def test_when_chat_agent_talks_to_a_real_http_endpoint_then_it_completes(tmp_path) -> None:
    with given() as context:
        server = ThreadingHTTPServer(("127.0.0.1", 0), _CompletionsHandler)
        threading.Thread(target=server.serve_forever, daemon=True).start()
        base_url = f"http://127.0.0.1:{server.server_address[1]}/v1"
        context.agent = ChatAgent("local", base_url=base_url, model="local-model")

    with when():
        context.result = context.agent.solve(_task(tmp_path))
        server.shutdown()

    with then():
        assert_that(context.result.status, equal_to("completed"))
        assert_that(context.result.log["exchanges"][0]["usage"], equal_to({"total_tokens": 42}))


def test_when_manual_agent_has_no_workbook_yet_then_it_waits_for_one(tmp_path) -> None:
    with given() as context:
        context.task = _task(tmp_path)
        context.agent = ManualAgent("person")

    with when():
        context.first = context.agent.solve(context.task)
        (context.task.workspace / OUTPUT_NAME).write_bytes(b"placeholder")
        context.second = context.agent.solve(context.task)

    with then():
        assert_that(context.first.status, equal_to("pending"))
        assert_that(context.first.output_workbook, none())
        assert_that((context.task.workspace / TASK_NAME).exists(), equal_to(True))
        assert_that(context.second.status, equal_to("completed"))


def test_when_task_has_custom_markdown_then_agents_use_it_verbatim(tmp_path) -> None:
    with given() as context:
        context.task = AgentTask(
            "custom",
            "ignored",
            _task(tmp_path).input_workbook,
            tmp_path / "workspace",
            task_markdown="# Custom\n\nSave output.xlsx.",
        )

    with when():
        context.text = task_text(context.task)

    with then():
        assert_that(context.text, equal_to("# Custom\n\nSave output.xlsx."))


def test_when_config_is_loaded_then_agents_of_each_type_are_built(tmp_path) -> None:
    with given() as context:
        context.config = tmp_path / "agents.toml"
        context.config.write_text(
            '[agents.coder]\ntype = "cli"\ncommand = ["coder", "{prompt}"]\n'
            '[agents.chat]\ntype = "chat"\nbase_url = "http://x/v1"\nmodel = "m"\n'
            '[agents.person]\ntype = "manual"\n',
            encoding="utf-8",
        )

    with when():
        context.agents = load_agents(context.config)

    with then():
        assert_that([agent.kind for agent in context.agents], equal_to(["cli", "chat", "manual"]))
        with pytest.raises(ValueError, match="unknown agent"):
            load_agents(context.config, ["nobody"])


@pytest.mark.parametrize(("program", "ready"), [(sys.executable, True), ("no-such-agent", False)])
def test_when_cli_agent_is_checked_then_its_executable_decides(program, ready) -> None:
    with given() as context:
        context.settings = {"type": "cli", "command": [program, "{prompt}"], "setup": "install"}

    with when():
        context.status = check_agent(context.settings)

    with then():
        assert_that(context.status.ready, equal_to(ready))
        assert_that(context.status.setup, equal_to("install"))
        if not ready:
            assert_that(context.status.detail, equal_to("no-such-agent not found"))


@pytest.mark.parametrize(("key", "ready"), [("secret", True), (None, False)])
def test_when_chat_agent_is_checked_then_its_api_key_decides(monkeypatch, key, ready) -> None:
    with given() as context:
        if key:
            monkeypatch.setenv("TICKMARK_TEST_KEY", key)
        else:
            monkeypatch.delenv("TICKMARK_TEST_KEY", raising=False)
        context.settings = {
            "type": "chat",
            "base_url": "http://localhost:1/v1",
            "model": "m",
            "api_key_env": "TICKMARK_TEST_KEY",
        }

    with when():
        context.status = check_agent(context.settings)

    with then():
        assert_that(context.status.ready, equal_to(ready))
        if not ready:
            assert_that(context.status.detail, equal_to("TICKMARK_TEST_KEY is not set"))


def test_when_shipped_config_is_read_then_every_agent_has_setup_steps() -> None:
    with when():
        context_table = load_agent_table(CONFIG)

    with then():
        missing = [name for name, settings in context_table.items() if not settings.get("setup")]
        assert_that(missing, equal_to([]))
