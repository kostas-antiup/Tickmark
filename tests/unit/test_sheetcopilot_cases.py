from pathlib import Path

import openpyxl
from hamcrest import assert_that, contains_exactly, equal_to
from tests.shared.givenpy import given, then, when

from tickmark.sheetcopilot_cases import (
    load_sheetcopilot_cases,
    sheetcopilot_prompt,
)


def _write_dataset(root: Path) -> None:
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Sheet Name", "No.", "Context", "Instructions", "Categories", "Atomic actions"])
    sheet.append(
        [
            "BoomerangSales",
            1,
            "A sales workbook.",
            "Calculate the revenue column.",
            "Formula, Entry & manipulation",
            "Write, AutoFill",
        ]
    )
    book.save(root / "dataset_20Samples.xlsx")

    source_dir = root / "task_sheets"
    source_dir.mkdir()
    source = openpyxl.Workbook()
    source.active["A1"] = "Product"
    source.save(source_dir / "BoomerangSales.xlsx")

    reference_dir = root / "task_sheet_answers_v2" / "BoomerangSales" / "1_BoomerangSales"
    reference_dir.mkdir(parents=True)
    reference = openpyxl.Workbook()
    reference.active["A1"] = "Revenue"
    reference.save(reference_dir / "1_BoomerangSales_gt1.xlsx")
    (reference_dir / "1_BoomerangSales_gt1_check.yaml").write_text(
        """
check_board:
  '1':
    cells:
      values: true
      formatting: false
      hyperlink: false
required_APIs:
- Write
- AutoFill
""".lstrip(),
        encoding="utf-8",
    )


def test_when_sheetcopilot_dataset_is_loaded_then_cases_follow_upstream_layout(tmp_path) -> None:
    with given() as context:
        _write_dataset(tmp_path)

    with when():
        context.cases = load_sheetcopilot_cases(tmp_path, case_ids=["1_BoomerangSales"])
        context.case = context.cases[0]

    with then():
        assert_that([case.case_id for case in context.cases], contains_exactly("1_BoomerangSales"))
        assert_that(context.case.input_workbook.name, equal_to("BoomerangSales.xlsx"))
        assert_that(context.case.categories, contains_exactly("Formula", "Entry & manipulation"))
        assert_that(context.case.atomic_actions, contains_exactly("Write", "AutoFill"))
        assert_that(context.case.references[0].required_apis, contains_exactly("Write", "AutoFill"))
        assert_that(
            context.case.references[0].check_board["1"]["cells"]["values"],
            equal_to(True),
        )
        assert_that(
            sheetcopilot_prompt(context.case),
            equal_to("Context: A sales workbook.\n\nTask: Calculate the revenue column."),
        )
