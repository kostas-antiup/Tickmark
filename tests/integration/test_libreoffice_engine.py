"""Real LibreOffice recalculation; skipped unless ``soffice`` is installed."""

from pathlib import Path

import openpyxl
import pytest
from hamcrest import assert_that, close_to, contains_exactly, equal_to
from tests.shared.givenpy import given, then, when

from tickmark.cases import load_case
from tickmark.evaluator import evaluate_workbook
from tickmark.recalculation import LibreOfficeRecalculator, find_soffice

pytestmark = [
    pytest.mark.libreoffice,
    pytest.mark.skipif(find_soffice() is None, reason="LibreOffice (soffice) not installed"),
]

CASE_14_07 = Path(__file__).resolve().parents[2] / "data" / "cases" / "14_07"


def _completed_14_07(path: Path, overrides: dict[str, object]) -> Path:
    """Write 14_07 as an agent would: golden formulas into input.xlsx via openpyxl."""

    golden = openpyxl.load_workbook(CASE_14_07 / "golden.xlsx")["RentRoll"]
    book = openpyxl.load_workbook(CASE_14_07 / "input.xlsx")
    sheet = book["RentRoll"]
    for row in golden["C17:G20"]:
        for cell in row:
            sheet[cell.coordinate] = overrides.get(cell.coordinate, cell.value)
    book.save(path)
    return path


def test_when_workbook_has_no_saved_values_then_libreoffice_calculates_them(tmp_path) -> None:
    with given() as context:
        book = openpyxl.Workbook()
        sheet = book.worksheets[0]
        sheet.title = "Model"
        sheet["B2"], sheet["B3"], sheet["B10"] = 100, 0.2, "=B2*(1+B3)"
        context.path = tmp_path / "model.xlsx"
        book.save(context.path)

    with when(), LibreOfficeRecalculator() as engine:
        context.base = engine.evaluate(context.path, ["Model!B10"])
        context.changed = engine.evaluate(context.path, ["Model!B10"], {"Model!B2": 110})

    with then():
        assert_that(context.base["Model!B10"], close_to(120, 1e-9))
        assert_that(context.changed["Model!B10"], close_to(132, 1e-9))


def test_when_agent_writes_correct_formulas_then_libreoffice_grades_them_as_passing(
    tmp_path,
) -> None:
    with given() as context:
        context.case = load_case(CASE_14_07)
        context.workbook = _completed_14_07(tmp_path / "correct.xlsx", {})

    with when(), LibreOfficeRecalculator() as engine:
        context.report = evaluate_workbook(context.case, context.workbook, engine)

    with then():
        assert_that(context.report.remarks, equal_to(()))
        assert_that(context.report.passed, equal_to(True))


def test_when_one_column_is_hardcoded_into_the_total_then_perturbation_catches_it(
    tmp_path,
) -> None:
    with given() as context:
        context.case = load_case(CASE_14_07)
        context.workbook = _completed_14_07(
            tmp_path / "partial.xlsx", {"G20": "=C20+D20+E20+1161600"}
        )

    with when(), LibreOfficeRecalculator() as engine:
        context.report = evaluate_workbook(context.case, context.workbook, engine)

    with then():
        assert_that(context.report.correct, equal_to(True))
        assert_that(context.report.scores["traceability"], equal_to(1.0))
        assert_that(
            context.report.checks["perturbation_check"].cells, contains_exactly("RentRoll!G20")
        )
        assert_that(context.report.passed, equal_to(False))
